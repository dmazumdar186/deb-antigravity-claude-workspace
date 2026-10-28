"""build_answer_sheet.py — builds the fillable ASSET answer sheet (PDF with AcroForm fields + Word twin)
from the 30 questions on site/questions/index.html, so the sheet never drifts from the page.

usage: python3 -I research/answer_sheet/build_answer_sheet.py <deliverable dir>
outputs: <D>/site/questions/ASSET_Answer_Sheet.pdf and .docx
"""
import html
import re
import sys
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Flowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer)

D = Path(sys.argv[1]).resolve()
SRC = D / "site/questions/index.html"
OUT = D / "site/questions"
DATE = "28 October 2026"
DEADLINE = "Friday 6 November 2026"
EMAIL = "debanjan@prodcraft.fyi"
INK, OR_INK, OR, MUTED = colors.HexColor("#0A1F30"), colors.HexColor("#9A4705"), colors.HexColor("#EF760D"), colors.HexColor("#4A5A68")
LABELS = ["Dates", "Budget", "Approvals", "Current provider", "Accounts", "Domain", "Analytics, trackers", "Selectour",
          "Texts, pages", "Brand", "Photos, logos", "Small-business offer", "Assistance", "English", "Enquiry recipients",
          "CRM", "Congress file", "GM&E codes", "Congress approval", "Fare wording", "Past congresses", "Quote request",
          "Client login", "Wondermiles", "AI option", "Staff login", "Data location", "Admin users", "Monthly changes",
          "Support hours"]
INSTR = "Write OK to accept our assumption, or your answer."


def text(fragment):
    return html.unescape(re.sub(r"<[^>]+>", "", fragment)).strip()


def parse():
    s = SRC.read_text(encoding="utf-8")
    groups = []
    for g, body in re.findall(r'<h3>([^<]+)</h3>\s*<ol class="qs">(.*?)</ol>', s, re.S):
        items = []
        for n, inner in re.findall(r'<li id="q(\d+)"><b>\d+</b><span>(.*?)</span></li>', body, re.S):
            first = '<span class="af">Answer first</span>' in inner
            inner = inner.replace('<span class="af">Answer first</span>', "")
            q, d = inner.split('<span class="qa">', 1)
            d = re.sub(r"<i>.*?</i>", "", d, count=1)
            items.append({"n": int(n), "label": LABELS[int(n) - 1], "q": text(q), "d": text(d), "first": first})
        groups.append((html.unescape(g), items))
    qs = [q for _, it in groups for q in it]
    assert [q["n"] for q in qs] == list(range(1, 31)), "expected q1..q30 in order"
    assert sum(q["first"] for q in qs) == 8, "expected 8 answer-first questions"
    return groups


def fonts():
    base = "/usr/share/fonts/truetype/crosextra/"
    for name, f in [("Body", "Carlito-Regular.ttf"), ("Body-B", "Carlito-Bold.ttf"), ("Body-I", "Carlito-Italic.ttf"),
                    ("Body-BI", "Carlito-BoldItalic.ttf")]:
        pdfmetrics.registerFont(TTFont(name, base + f))
    pdfmetrics.registerFontFamily("Body", normal="Body", bold="Body-B", italic="Body-I", boldItalic="Body-BI")


class Field(Flowable):
    """A labelled multiline AcroForm text field that flows with the text."""

    def __init__(self, name, tooltip, height, label=None, width=None, multiline=True):
        super().__init__()
        self.name, self.tooltip, self.h, self.label, self.w, self.multiline = name, tooltip, height, label, width, multiline

    def wrap(self, aw, ah):
        self.width = self.w or aw
        return self.width, self.h + (12 if self.label else 0)

    def draw(self):
        c = self.canv
        if self.label:
            c.setFont("Body-B", 8.5); c.setFillColor(MUTED); c.drawString(0, self.h + 3, self.label.upper())
        c.acroForm.textfield(name=self.name, tooltip=self.tooltip, x=0, y=0, width=self.width, height=self.h,
                             relative=True, borderColor=colors.HexColor("#9AA5AE"), fillColor=colors.HexColor("#FBFAF6"),
                             textColor=INK, borderWidth=0.8, fontName="Helvetica", fontSize=10 if self.multiline else 11,
                             fieldFlags="multiline" if self.multiline else "", maxlen=4000, forceBorder=True)


class Row(Flowable):
    """Two single-line fields side by side (Answered by: name, role)."""

    def __init__(self, fields):
        super().__init__(); self.fields = fields

    def wrap(self, aw, ah):
        self.width = aw; return aw, 32

    def draw(self):
        gap = 10; w = (self.width - gap) / 2
        for i, (name, label) in enumerate(self.fields):
            self.canv.saveState(); self.canv.translate(i * (w + gap), 0)
            f = Field(name, label, 20, label=label, width=w, multiline=False); f.canv = self.canv; f.wrap(w, 40); f.draw()
            self.canv.restoreState()


def build_pdf(groups):
    fonts()
    st = lambda **k: ParagraphStyle("x", **{"fontName": "Body", "fontSize": 10, "leading": 13, "textColor": INK, "alignment": TA_LEFT, **k})
    s_brand, s_h1, s_p = st(fontName="Body-B", fontSize=9, textColor=OR_INK), st(fontName="Body-B", fontSize=20, leading=24), st()
    s_inst = st(fontName="Body-B", fontSize=11.5, leading=15)
    s_grp = st(fontName="Body-B", fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=4, textColor=OR_INK)
    s_qh, s_q = st(fontName="Body-B", fontSize=11, leading=14), st(fontSize=10.5, leading=13.5)
    s_d = st(fontName="Body-I", fontSize=9.5, leading=12.5, textColor=MUTED)
    out = OUT / "ASSET_Answer_Sheet.pdf"

    def chrome(c, doc):
        c.saveState(); c.setFont("Body", 8); c.setFillColor(MUTED)
        c.drawString(18 * mm, 10 * mm, "ASSET Voyages · answer sheet · ProdCraft")
        c.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Return as one PDF to {EMAIL} by 6 November 2026 · page {doc.page}")
        c.setStrokeColor(OR); c.setLineWidth(2); c.line(18 * mm, A4[1] - 12 * mm, A4[0] - 18 * mm, A4[1] - 12 * mm)
        c.restoreState()

    doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm,
                            bottomMargin=17 * mm, title="ASSET Voyages – answer sheet (30 questions)",
                            author="ProdCraft", subject="Answers to ProdCraft's 30 questions", creator="ProdCraft")
    af = ", ".join(str(q["n"]) for _, it in groups for q in it if q["first"])
    story = [Paragraph("ASSET VOYAGES &nbsp;·&nbsp; PRODCRAFT", s_brand), Spacer(1, 4),
             Paragraph("Answer sheet: our 30 questions", s_h1), Spacer(1, 3),
             Paragraph(f"Website redesign · questions sent {DATE}", st(textColor=MUTED)), Spacer(1, 10),
             Paragraph("Answered by", st(fontName="Body-B", fontSize=11)), Spacer(1, 4),
             Row([("answered_by_name", "Name"), ("answered_by_role", "Role")]), Spacer(1, 12),
             Paragraph(INSTR, s_inst), Spacer(1, 4),
             Paragraph(f"Please return your answers as one PDF to <b>{EMAIL}</b> by <b>{DEADLINE}</b>. Fill in this PDF "
                       "directly, or use the Word version and save it as PDF. Short answers are fine. Anything not answered "
                       "by then follows the assumption shown under the question.", s_p), Spacer(1, 4),
             Paragraph(f"<font color='#9A4705'><b>ANSWER FIRST</b></font> marks the eight questions to start with if "
                       f"time is short: {af}.", s_p), Spacer(1, 2)]
    for g, items in groups:
        story.append(Paragraph(g, s_grp))
        for q in items:
            tag = " &nbsp;<font color='#9A4705' size='8.5'><b>ANSWER FIRST</b></font>" if q["first"] else ""
            label = html.escape(q["label"])
            block = [Paragraph(f"<b>{q['n']}</b> · {label}{tag}", s_qh), Spacer(1, 2),
                     Paragraph(html.escape(q["q"]), s_q), Spacer(1, 3),
                     Paragraph("Our assumption if we don’t hear back: " + html.escape(q["d"]), s_d), Spacer(1, 4),
                     Field(f"q{q['n']:02d}_answer", f"Question {q['n']} ({q['label']}): {INSTR}", 58), Spacer(1, 12)]
            story.append(KeepTogether(block))
    doc.build(story, onFirstPage=chrome, onLaterPages=chrome)
    return out


def build_docx(groups):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    ink, orink, muted = RGBColor(0x0A, 0x1F, 0x30), RGBColor(0x9A, 0x47, 0x05), RGBColor(0x4A, 0x5A, 0x68)
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Mm(18); s.top_margin = s.bottom_margin = Mm(16)
    normal = d.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(10.5); normal.font.color.rgb = ink
    d.core_properties.title = "ASSET Voyages – answer sheet (30 questions)"; d.core_properties.author = "ProdCraft"

    def para(txt="", bold=False, italic=False, size=None, color=None, after=4):
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
        if txt:
            r = p.add_run(txt); r.bold, r.italic = bold, italic
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
        return p

    def box(label=None, height_lines=3):
        t = d.add_table(rows=1, cols=1); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT
        cell = t.rows[0].cells[0]; cell.paragraphs[0].text = ""
        for _ in range(height_lines - 1):
            cell.add_paragraph("")
        tc = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "FBFAF6"); tc.append(shd)
        para(after=6)

    para("ASSET VOYAGES  ·  PRODCRAFT", bold=True, size=9, color=orink, after=2)
    para("Answer sheet: our 30 questions", bold=True, size=20, after=2)
    para(f"Website redesign · questions sent {DATE}", size=10, color=muted, after=10)
    para("Answered by", bold=True, size=11, after=3)
    t = d.add_table(rows=2, cols=2); t.style = "Table Grid"
    t.cell(0, 0).text, t.cell(0, 1).text = "Name", "Role"
    for c in t.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True; c.paragraphs[0].runs[0].font.size = Pt(9)
    para(after=6)
    para(INSTR, bold=True, size=12, after=3)
    para(f"Please return your answers as one PDF to {EMAIL} by {DEADLINE}. Type your answers in the boxes below, "
         "then save this file as PDF. Short answers are fine. Anything not answered by then follows the assumption shown "
         "under the question.", after=3)
    af = ", ".join(str(q["n"]) for _, it in groups for q in it if q["first"])
    p = para(after=6); r = p.add_run("ANSWER FIRST"); r.bold = True; r.font.color.rgb = orink; r.font.size = Pt(9)
    p.add_run(f" marks the eight questions to start with if time is short: {af}.")
    for g, items in groups:
        h = para(g, bold=True, size=13, color=orink, after=4); h.paragraph_format.space_before = Pt(10)
        for q in items:
            p = para(after=2); p.paragraph_format.keep_with_next = True
            r = p.add_run(f"{q['n']} · {q['label']}"); r.bold = True; r.font.size = Pt(11.5)
            if q["first"]:
                r = p.add_run("   ANSWER FIRST"); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = orink
            p = para(q["q"], after=2); p.paragraph_format.keep_with_next = True
            p = para("Our assumption if we don’t hear back: " + q["d"], italic=True, size=9.5, color=muted, after=3)
            p.paragraph_format.keep_with_next = True
            box()
    out = OUT / "ASSET_Answer_Sheet.docx"; d.save(str(out)); return out


if __name__ == "__main__":
    gs = parse()
    print(build_pdf(gs)); print(build_docx(gs))
