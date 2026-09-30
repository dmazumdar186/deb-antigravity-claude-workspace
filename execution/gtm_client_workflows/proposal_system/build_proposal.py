"""
build_proposal.py
description: Render a proposal JSON (proposals/<slug>.json) into an editable DOCX (Google-Docs friendly, plain styling) and a tracked HTML page for the proposal-tracker Worker.
inputs: proposals/<slug>.json — {slug,title,version,date,parties,blocks[]} where blocks are {"h1"|"h2"|"h3"|"p"|"bullets"|"numbered"|"table"|"note"|"signature"|"pagebreak"}; optional --pixel-url for the HTML open pixel.
outputs: out/<slug>/<slug>.docx, out/<slug>/<slug>.html, out/<slug>/<slug>.md (paste-ready plain text). Exit 1 on schema errors.
"""
from __future__ import annotations

import argparse
import html
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
BLOCK_TYPES = {"h1", "h2", "h3", "p", "bullets", "numbered", "table", "note", "signature", "pagebreak", "kv"}


def load(slug: str) -> dict:
    path = HERE / "proposals" / f"{slug}.json"
    if not path.exists():
        sys.exit(f"missing {path}")
    data = json.loads(path.read_text(encoding="utf-8"))
    for k in ("slug", "title", "version", "date", "parties", "blocks"):
        if k not in data:
            sys.exit(f"schema: missing top-level key '{k}'")
    for i, b in enumerate(data["blocks"]):
        if len(b) != 1 or next(iter(b)) not in BLOCK_TYPES:
            sys.exit(f"schema: block {i} must have exactly one key from {sorted(BLOCK_TYPES)}: {b}")
    return data


# ---------- inline markup: **bold** only (keeps the source readable) ----------
_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _runs(par, text: str, size: int | None = None):
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            if size:
                r.font.size = Pt(size)
        r = par.add_run(m.group(1))
        r.bold = True
        if size:
            r.font.size = Pt(size)
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        if size:
            r.font.size = Pt(size)


def _inline_html(text: str) -> str:
    return _BOLD.sub(r"<strong>\1</strong>", html.escape(text, quote=False))


# ---------- DOCX ----------
def build_docx(d: dict, out: Path) -> None:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    for lvl, sz in ((1, 20), (2, 15), (3, 12)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = "Arial"
        h.font.size = Pt(sz)
        h.font.color.rgb = RGBColor(0, 0, 0)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _runs(t, f"**{d['title']}**", 22)
    meta = doc.add_paragraph()
    _runs(meta, f"Version {d['version']} · {d['date']}" + (f" · Valid until {d['valid_until']}" if d.get("valid_until") else ""))
    if d.get("subtitle"):
        doc.add_paragraph(d["subtitle"])

    for b in d["blocks"]:
        k, v = next(iter(b.items()))
        if k == "h1":
            doc.add_heading(v, level=1)
        elif k == "h2":
            doc.add_heading(v, level=2)
        elif k == "h3":
            doc.add_heading(v, level=3)
        elif k == "p":
            _runs(doc.add_paragraph(), v)
        elif k == "note":
            p = doc.add_paragraph()
            r = p.add_run(v)
            r.italic = True
            r.font.size = Pt(9.5)
        elif k in ("bullets", "numbered"):
            style = "List Bullet" if k == "bullets" else "List Number"
            for item in v:
                _runs(doc.add_paragraph(style=style), item)
        elif k == "kv":
            tbl = doc.add_table(rows=0, cols=2)
            tbl.style = "Table Grid"
            for key, val in v:
                row = tbl.add_row().cells
                _runs(row[0].paragraphs[0], f"**{key}**")
                _runs(row[1].paragraphs[0], val)
        elif k == "table":
            hdr, rows = v["header"], v["rows"]
            tbl = doc.add_table(rows=1, cols=len(hdr))
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
            for i, h in enumerate(hdr):
                _runs(tbl.rows[0].cells[i].paragraphs[0], f"**{h}**")
            for r in rows:
                if len(r) != len(hdr):
                    sys.exit(f"table row width mismatch: {r}")
                cells = tbl.add_row().cells
                for i, c in enumerate(r):
                    _runs(cells[i].paragraphs[0], str(c))
            if v.get("caption"):
                p = doc.add_paragraph()
                rr = p.add_run(v["caption"])
                rr.italic = True
                rr.font.size = Pt(9.5)
        elif k == "pagebreak":
            doc.add_page_break()
        elif k == "signature":
            # two-column signature layout: left = provider, right = client
            tbl2 = doc.add_table(rows=6, cols=2)
            tbl2.style = "Table Grid"
            for ci, party in enumerate(v):
                cells = [tbl2.rows[r].cells[ci] for r in range(6)]
                _runs(cells[0].paragraphs[0], f"**{party['for']}**")
                _runs(cells[1].paragraphs[0], f"Name: {party['name']}")
                _runs(cells[2].paragraphs[0], f"Title: {party['title']}")
                _runs(cells[3].paragraphs[0], "Signature: ______________________________")
                _runs(cells[4].paragraphs[0], "Date: ____ / ____ / ________")
                _runs(cells[5].paragraphs[0], f"Email: {party['email']}")

    doc.save(out)


# ---------- HTML (tracked web copy) ----------
# Mirrors the tracking contract of the reference proposal-system (api/track.js):
# first open per browser in 24 h -> "opened", later loads -> "viewed".
TRACK_JS = """<script>
const TRACK_URL = "__TRACK_URL__", TRACK_SLUG = "__SLUG__", TRACK_CLIENT = "__CLIENT__";
function track(event, extra) { try { fetch(TRACK_URL, { method: 'POST', headers: { 'Content-Type': 'application/json' },
  body: JSON.stringify({ slug: TRACK_SLUG, client: TRACK_CLIENT, event, extra }), keepalive: true }).catch(() => {}); } catch (e) {} }
(function () { try { const key = 'proposal_opened_' + TRACK_SLUG; const last = parseInt(localStorage.getItem(key) || '0', 10); const now = Date.now();
  if (!last || now - last > 24 * 60 * 60 * 1000) { track('opened'); localStorage.setItem(key, String(now)); } else { track('viewed'); } } catch (e) { track('opened'); } })();
</script>"""

CSS = """body{font-family:Arial,Helvetica,sans-serif;max-width:820px;margin:40px auto;padding:0 20px;color:#111;line-height:1.5}
h1{font-size:26px;margin:28px 0 8px}h2{font-size:20px;margin:26px 0 8px}h3{font-size:16px;margin:18px 0 6px}
table{border-collapse:collapse;width:100%;margin:10px 0 16px;font-size:14px}td,th{border:1px solid #999;padding:6px 8px;vertical-align:top;text-align:left}
.note{font-style:italic;font-size:13px;color:#444}.meta{color:#555}.sig td{height:40px}"""


def build_html(d: dict, out: Path, pixel_url: str | None, track_url: str | None = None) -> None:
    parts = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='robots' content='noindex,nofollow'>"
             f"<meta name='viewport' content='width=device-width,initial-scale=1'><title>{html.escape(d['title'])}</title><style>{CSS}</style></head><body>",
             f"<h1>{html.escape(d['title'])}</h1>",
             f"<p class='meta'>Version {html.escape(str(d['version']))} · {html.escape(d['date'])}"
             + (f" · Valid until {html.escape(d['valid_until'])}" if d.get("valid_until") else "") + "</p>"]
    if d.get("subtitle"):
        parts.append(f"<p>{_inline_html(d['subtitle'])}</p>")
    for b in d["blocks"]:
        k, v = next(iter(b.items()))
        if k in ("h1", "h2", "h3"):
            parts.append(f"<{k}>{_inline_html(v)}</{k}>")
        elif k == "p":
            parts.append(f"<p>{_inline_html(v)}</p>")
        elif k == "note":
            parts.append(f"<p class='note'>{_inline_html(v)}</p>")
        elif k in ("bullets", "numbered"):
            tag = "ul" if k == "bullets" else "ol"
            parts.append(f"<{tag}>" + "".join(f"<li>{_inline_html(i)}</li>" for i in v) + f"</{tag}>")
        elif k == "kv":
            parts.append("<table>" + "".join(f"<tr><th>{_inline_html(a)}</th><td>{_inline_html(b2)}</td></tr>" for a, b2 in v) + "</table>")
        elif k == "table":
            parts.append("<table><tr>" + "".join(f"<th>{_inline_html(h)}</th>" for h in v["header"]) + "</tr>"
                         + "".join("<tr>" + "".join(f"<td>{_inline_html(str(c))}</td>" for c in r) + "</tr>" for r in v["rows"]) + "</table>"
                         + (f"<p class='note'>{_inline_html(v['caption'])}</p>" if v.get("caption") else ""))
        elif k == "pagebreak":
            parts.append("<hr>")
        elif k == "signature":
            parts.append("<table class='sig'><tr>" + "".join(f"<th>{html.escape(p['for'])}</th>" for p in v) + "</tr>"
                         + "".join("<tr>" + "".join(f"<td>{lab}: {html.escape(str(p.get(key, '')))}</td>" for p in v) + "</tr>"
                                   for lab, key in (("Name", "name"), ("Title", "title"), ("Email", "email")))
                         + "<tr>" + "".join("<td>Signature: ______________________</td>" for _ in v) + "</tr>"
                         + "<tr>" + "".join("<td>Date: ____ / ____ / ________</td>" for _ in v) + "</tr></table>")
    if pixel_url:
        parts.append(f"<img src='{html.escape(pixel_url)}' width='1' height='1' alt='' style='position:absolute;left:-9999px'>")
    if track_url:
        parts.append(TRACK_JS.replace("__TRACK_URL__", track_url).replace("__SLUG__", d["slug"]).replace("__CLIENT__", html.escape(d["parties"]["client"]["company"])))
    parts.append("</body></html>")
    out.write_text("\n".join(parts), encoding="utf-8")


# ---------- Markdown (paste-ready) ----------
def build_md(d: dict, out: Path) -> None:
    L = [f"# {d['title']}", f"Version {d['version']} · {d['date']}" + (f" · Valid until {d['valid_until']}" if d.get("valid_until") else ""), ""]
    if d.get("subtitle"):
        L += [d["subtitle"], ""]
    for b in d["blocks"]:
        k, v = next(iter(b.items()))
        if k == "h1":
            L += [f"## {v}", ""]
        elif k == "h2":
            L += [f"### {v}", ""]
        elif k == "h3":
            L += [f"#### {v}", ""]
        elif k in ("p", "note"):
            L += [v, ""]
        elif k == "bullets":
            L += [f"- {i}" for i in v] + [""]
        elif k == "numbered":
            L += [f"{n}. {i}" for n, i in enumerate(v, 1)] + [""]
        elif k == "kv":
            L += ["| | |", "|---|---|"] + [f"| **{a}** | {b2} |" for a, b2 in v] + [""]
        elif k == "table":
            L += ["| " + " | ".join(v["header"]) + " |", "|" + "---|" * len(v["header"])]
            L += ["| " + " | ".join(str(c) for c in r) + " |" for r in v["rows"]] + ([v["caption"], ""] if v.get("caption") else [""])
        elif k == "pagebreak":
            L += ["---", ""]
        elif k == "signature":
            for p in v:
                L += [f"**{p['for']}**", f"Name: {p['name']}", f"Title: {p['title']}", f"Email: {p['email']}", "Signature: ____________", "Date: ____", ""]
    out.write_text("\n".join(L), encoding="utf-8")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("slug")
    ap.add_argument("--pixel-url", default=None, help="open-pixel URL to embed in the HTML copy, e.g. https://<worker>/o/<slug>.gif")
    ap.add_argument("--track-url", default=None, help="POST endpoint for opened/viewed events, e.g. https://<worker>/api/track")
    a = ap.parse_args()
    d = load(a.slug)
    out_dir = HERE / "out" / a.slug
    out_dir.mkdir(parents=True, exist_ok=True)
    build_docx(d, out_dir / f"{a.slug}.docx")
    build_html(d, out_dir / f"{a.slug}.html", a.pixel_url, a.track_url)
    build_md(d, out_dir / f"{a.slug}.md")
    n_tables = sum(1 for b in d["blocks"] if "table" in b)
    print(f"built {out_dir}: docx, html, md · {len(d['blocks'])} blocks, {n_tables} tables")
    return 0


if __name__ == "__main__":
    sys.exit(main())
