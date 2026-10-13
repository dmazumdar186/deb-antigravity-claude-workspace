#!/usr/bin/env python3
"""
instructor_sheet_export.py
description: Turn fr_instructor_finder.py output into an outreach workbook: one tab per city,
             columns First name | Last name | Email | Website exists | Website | Found on.
             Writes .xlsx (uploadable to Drive, converts to a Google Sheet with the tabs kept).
inputs:  --input .tmp/fr_instructors.json  --out .tmp/yoga_teacher.xlsx  [--include-offdomain]
outputs: the .xlsx plus a per-city summary on stdout
usage:   python3 execution/lead_sourcing/instructor_sheet_export.py --input .tmp/yoga_fr.json --out .tmp/yoga_teacher.xlsx
"""
import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

HEADERS = ["First name", "Last name", "Email", "Website exists", "Website", "Found on"]
ORG_WORDS = re.compile(r"\b(yoga|studio|centre|center|school|école|ecole|pilates|club|association|salle|espace|atelier|maison|cours|institut|gym|wellness|zen|shala|academy|académie)\b", re.I)
SOCIAL = ("instagram.com", "facebook.com", "linktr.ee", "linkedin.com")


def split_name(lead: dict) -> tuple[str, str]:
    if lead.get("first_name"):
        return lead["first_name"].strip(), (lead.get("last_name") or "").strip()
    name = (lead.get("name") or "").strip()
    if ORG_WORDS.search(name) or len(name.split()) not in (2, 3):
        return name, ""  # a studio/brand name: keep it whole in the first column
    parts = name.split()
    return parts[0], " ".join(parts[1:])


def website_status(lead: dict) -> tuple[str, str]:
    site = (lead.get("website") or "").strip()
    if not site:
        return "no", ""
    if any(s in site for s in SOCIAL):
        return "no", site  # social page only: sells the site, keep the link
    if lead.get("website_source") == "serper_enrich":
        return "not sure", site  # name-based guess, verify before pitching
    return "yes", site


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", default=".tmp/fr_instructors.json")
    ap.add_argument("--out", default=".tmp/yoga_teacher.xlsx")
    ap.add_argument("--include-offdomain", action="store_true", help="keep emails whose domain differs from the site")
    a = ap.parse_args()

    leads = json.loads(Path(a.input).read_text(encoding="utf-8"))
    by_city: dict[str, list[list[str]]] = defaultdict(list)
    seen = set()
    for lead in leads:
        city = (lead.get("city") or "Unknown").strip().title()
        first, last = split_name(lead)
        key = (first.lower(), last.lower(), city)
        if key in seen:
            continue
        seen.add(key)
        email = lead.get("email") or ""
        # A gmail-style address on a site Google Maps lists for this teacher is theirs; the same address on a
        # site we only guessed from a first name is probably a directory's, so drop it unless asked.
        if (lead.get("email_source") == "website_offdomain" and lead.get("website_source") == "serper_enrich"
                and not a.include_offdomain):
            email = ""
        exists, site = website_status(lead)
        by_city[city].append([first, last, email, exists, site, lead.get("source", "")])

    wb = Workbook()
    wb.remove(wb.active)
    summary = Counter()
    for city in sorted(by_city):
        ws = wb.create_sheet(title=city[:31])
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True)
        rows = sorted(by_city[city], key=lambda r: (r[2] == "", r[3] != "no", r[0]))
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for col, width in zip("ABCDEF", (18, 18, 32, 14, 40, 12)):
            ws.column_dimensions[col].width = width
        summary[city] = len(rows)
        summary[city + " (email)"] = sum(1 for r in rows if r[2])
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    for city in sorted(by_city):
        print(f"{city:14s} {summary[city]:4d} rows  {summary[city + ' (email)']:4d} with email")
    print(f"wrote {out} ({sum(len(v) for v in by_city.values())} rows, {len(by_city)} tabs)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
